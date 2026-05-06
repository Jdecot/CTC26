import pandas as pd
from decimal import Decimal, InvalidOperation
import config
from pathlib import Path
from openpyxl.styles import PatternFill

# Liste des colonnes techniques ou de métadonnées à exclure des calculs d'actifs
TECHNICAL_COLUMNS = [
    'Date', 'refid', 'subtype', 'Type', 'Detected Type', 'platform', 'Platform',
    'Normalized Received Currency', 'Normalized Sent Currency', 'Normalized Fee Currency',
    'Sent Currency', 'Sent Amount', 'Sent Net Worth', 
    'Received Currency', 'Received Amount', 'Received Net Worth', 
    'Fee Currency', 'Fee Amount', 'Fee Net Worth', 'PTA_sell_ratio',
    'Balance', 'is_taxable_event', 'Money_movement', 'wallet_value_before', 'wallet_value_after',
    'wallet_value_eur', 'wallet_value_eur_m1'
]

# Liste des devises Fiat (ou assimilées) à exclure des colonnes "crypto"
FIAT_CURRENCIES = ['EUR', 'USD', 'ZEUR', 'ZUSD']

def get_crypto_columns(df, include_stables=True):
    """
    Identifie les colonnes d'actifs cryptos (celles qui contiennent les soldes).
    On exclut les colonnes de métadonnées, les devises fiat et optionnellement les stables.
    """
    exclude = TECHNICAL_COLUMNS + FIAT_CURRENCIES
    if not include_stables:
        exclude.extend(['USDT', 'USDC'])
    
    # On ne garde que les colonnes qui ne sont pas dans la liste d'exclusion 
    # et qui ne sont pas des colonnes de prix (se terminant par _price)
    return [col for col in df.columns if col not in exclude and not col.endswith('_price')]

# def convert_csv_to_excel(csv_input_path, excel_output_path):
#     """
#     Converts a single CSV file to an Excel file.
#     Ensures the destination directory exists before writing.
#     """
#     try:
#         # Create parent directory if it does not exist
#         Path(excel_output_path).parent.mkdir(parents=True, exist_ok=True)
        
#         # Read CSV data and export to Excel format
#         # Standard comma separator used by default
#         data_frame = pd.read_csv(csv_input_path, sep=',')
#         data_frame.to_excel(excel_output_path, index=False)
#         return True
#     except Exception as error:
#         print(f"  ✗ Error during conversion of {csv_input_path}: {error}")
#         return False
    


def convert_csv_to_excel(csv_input_path, excel_output_path):
    """
    Converts a single CSV file to an Excel file.
    Ensures the destination directory exists before writing.
    Applies yellow header color and freezes the top row.
    """
    try:
        # Create parent directory if it does not exist
        Path(excel_output_path).parent.mkdir(parents=True, exist_ok=True)
        
        # Read CSV data
        data_frame = pd.read_csv(csv_input_path, sep=',')
        
        # Utilisation de ExcelWriter pour pouvoir manipuler le formatage
        with pd.ExcelWriter(excel_output_path, engine='openpyxl') as writer:
            data_frame.to_excel(writer, index=False, sheet_name='Sheet1')
            
            # On récupère les objets de la feuille de calcul
            workbook = writer.book
            worksheet = writer.sheets['Sheet1']
            
            # 1. COLORIER LA LIGNE 1 EN JAUNE
            # On définit le remplissage jaune (Code Hexa FFFF00)
            yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
            
            # On boucle sur chaque cellule de la première ligne
            # len(data_frame.columns) nous donne le nombre de colonnes
            for col_num in range(1, len(data_frame.columns) + 1):
                cell = worksheet.cell(row=1, column=col_num)
                cell.fill = yellow_fill
            
            # 2. FIGER LES VOLETS (Ligne supérieure)
            # Dans openpyxl, on définit la cellule qui sert de point d'ancrage. 
            # 'A2' fige tout ce qui est AU-DESSUS de la ligne 2.
            worksheet.freeze_panes = 'A2'
            
        return True
        
    except Exception as error:
        print(f"   ✗ Error during conversion of {csv_input_path}: {error}")
        return False


def show_holdings():
    """Lit la dernière ligne de account_situation.csv et affiche les holdings triés."""
    csv_path = config.FILE_ACCOUNT_SITUATION
    
    if not csv_path.exists():
        print(f"⚠️  Fichier non trouvé : {csv_path}")
        return
    
    df = pd.read_csv(csv_path, dtype=str)
    if df.empty:
        print("⚠️  Le fichier account_situation.csv est vide")
        return
    
    last_row = df.iloc[-1]
    
    crypto_cols = get_crypto_columns(df)

    holdings = {}
    for col in crypto_cols:
        try:
            raw_val = str(last_row[col]).strip()
            # On ignore les colonnes vides ou contenant 'nan' pour éviter les erreurs de tri
            if raw_val == "" or raw_val.lower() == "nan":
                continue
                
            val = Decimal(raw_val)
            # On ne garde que les nombres finis (pas NaN, pas Infini) et non nuls
            if not val.is_finite():
                continue
            if val != Decimal('0'):
                holdings[col] = val
        except (ValueError, TypeError, InvalidOperation):
            # Ignore les colonnes non numériques ou valeurs non convertibles
            pass
    
    sorted_holdings = dict(sorted(holdings.items(), key=lambda x: x[1], reverse=True))
    
    print("\n" + "="*40)
    print("📊 HOLDINGS (dernière ligne)")
    print("="*40)
    for crypto, qty in sorted_holdings.items():
        # Formate en décimal avec 18 chiffres max après la virgule, en supprimant les zéros inutiles à la fin
        print(f"  {crypto}: {qty:.18f}".rstrip('0').rstrip('.'))

def reorder_columns(df):
    """Réorganise les colonnes pour mettre les métadonnées et montants au début."""
    cols_prioritaires = [
        "Date", "Type", "subtype", "Detected Type", "platform", "refid",
        
        "Sent Amount", "Sent Currency", "Normalized Sent Currency",
        
        "Received Amount", "Received Currency", "Normalized Received Currency", 
        "Balance",
        
        "Fee Amount", "Fee Currency", "Normalized Fee Currency",
        
        "Sent Net Worth", "Received Net Worth", "Fee Net Worth",
        
        "is_taxable_event", 
        "wallet_value_before", "wallet_value_after",
        "PTA_sell_ratio"
    ]

    cols_existantes = [c for c in cols_prioritaires if c in df.columns]
    autres_cols = [c for c in df.columns if c not in cols_existantes]
    return df[cols_existantes + autres_cols]
